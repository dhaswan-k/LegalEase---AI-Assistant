"""DOCX export using python-docx."""

import io
import os
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from document.formatter import (
    DISCLAIMER,
    FOOTER_TEXT,
    classify_line,
    sanitize_text,
    split_terms,
    strip_markdown,
)

FONT = "Times New Roman"
NAVY = RGBColor(0x14, 0x28, 0x4B)
_WITNESS_RE = re.compile(r"IN WITNESS|SIGNATURE", re.I)


def _set_font(run, size=12, bold=False, italic=False, color=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def _add_field(paragraph, instruction):
    run = paragraph.add_run()
    _set_font(run, 9, color=RGBColor(0x55, 0x55, 0x55))
    for tag, text in (("begin", None), (None, instruction), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def _shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def _add_terms_table(doc, terms):
    items = split_terms(terms)
    if not items:
        return
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.keep_with_next = True
    _set_font(p.add_run("Schedule A: Summary of Agreed Terms"), 13, bold=True, color=NAVY)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for cell, label in zip(hdr, ("No.", "Terms & Conditions")):
        cell.text = ""
        _set_font(cell.paragraphs[0].add_run(label), 11, bold=True, color=RGBColor(255, 255, 255))
        _shade(cell, "14284B")
    for i, item in enumerate(items, 1):
        row = table.add_row().cells
        row[0].text = ""
        row[1].text = ""
        _set_font(row[0].paragraphs[0].add_run(str(i)), 11)
        row[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_font(row[1].paragraphs[0].add_run(sanitize_text(item)), 11)
    table.autofit = False
    table.columns[0].width = Cm(1.6)
    table.columns[1].width = Cm(14.4)
    for row in table.rows:
        row.cells[0].width = Cm(1.6)
        row.cells[1].width = Cm(14.4)
    doc.add_paragraph()


def _render_lines(doc, lines, title_done):
    for raw in lines:
        kind = classify_line(raw)
        if kind == "blank":
            continue
        text = strip_markdown(raw)
        if kind == "heading" and not title_done:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(14)
            _set_font(p.add_run(text.upper()), 16, bold=True, color=NAVY)
            title_done = True
        elif kind == "heading":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            _set_font(p.add_run(text), 12, bold=True)
        elif kind in ("sub", "bullet"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1.0)
            p.paragraph_format.space_after = Pt(4)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            if kind == "bullet":
                text = "• " + re.sub(r"^[-*]\s+", "", text)
            _set_font(p.add_run(text), 12)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.15
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            title_done = True
            _set_font(p.add_run(text), 12)
    return title_done


def format_docx(text, doc_type, terms=None, logo_path=None) -> bytes:
    """Build a formatted .docx and return its bytes."""
    text = sanitize_text(text)
    doc = Document()

    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2.5)
    section.top_margin, section.bottom_margin = Cm(2.2), Cm(2.2)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(12)

    # Logo (first page) - optional and failure-tolerant
    if logo_path and os.path.isfile(logo_path):
        try:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(logo_path, width=Cm(4.0))
        except Exception:
            pass

    lines = text.split("\n")
    # Insert terms table before the signature block when one exists
    split_idx = None
    if split_terms(terms):
        for i in range(len(lines) - 1, -1, -1):
            if _WITNESS_RE.search(lines[i]):
                split_idx = i
                break
    if split_idx is not None:
        done = _render_lines(doc, lines[:split_idx], False)
        _add_terms_table(doc, terms)
        _render_lines(doc, lines[split_idx:], done)
    else:
        _render_lines(doc, lines, False)
        _add_terms_table(doc, terms)

    # Disclaimer
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    _set_font(p.add_run(DISCLAIMER), 9, italic=True, color=RGBColor(0x55, 0x55, 0x55))

    # Footer with page numbers
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(fp.add_run(f"{FOOTER_TEXT}  |  Page "), 9, color=RGBColor(0x55, 0x55, 0x55))
    _add_field(fp, "PAGE")
    _set_font(fp.add_run(" of "), 9, color=RGBColor(0x55, 0x55, 0x55))
    _add_field(fp, "NUMPAGES")

    doc.core_properties.title = sanitize_text(doc_type) or "Legal Document"
    doc.core_properties.author = "LegalEase AI"

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
