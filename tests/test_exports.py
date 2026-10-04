"""Offline tests for sanitising, models and exporters (no Gemini key needed).

Run from project root:  python -m unittest discover -s tests -v
"""

import importlib.util
import io
import sys
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from document.docx_generator import format_docx
from document.formatter import format_html_preview, safe_filename, sanitize_text, split_terms
from document.txt_generator import generate_txt

LOGO = str(ROOT / "assets" / "logo.png")

DOC_TYPES = ["NDA", "Employment Contract", "Lease Agreement", "Freelance Contract", "Service Agreement"]
TERMS = "Payment within 30 days; Confidentiality must be maintained; Salary: ₹40,000/month"


def sample_doc(kind: str, paragraphs: int = 3) -> str:
    body = "\n\n".join(
        f"{i}. Clause {i}\nThe parties agree to the “terms” — including ₹ amounts, café, ‘quotes’ and 中文 text number {i}."
        for i in range(1, paragraphs + 1)
    )
    return (
        f"# {kind}\n\nThis {kind} is made effective 03 October 2026 between Jane Doe and TechNova Inc.\n\n{body}\n\n"
        "IN WITNESS WHEREOF, the parties have signed.\n\n______________\nJane Doe\n\n______________\nTechNova Inc."
    )


class SanitizeTests(unittest.TestCase):
    def test_smart_chars_and_blank_lines(self):
        out = sanitize_text("“Hi” — it’s\n\n\n\n\nok  ")
        self.assertEqual(out, '"Hi" - it\'s\n\nok')

    def test_none_and_bytes(self):
        self.assertEqual(sanitize_text(None), "")
        self.assertEqual(sanitize_text(b"abc"), "abc")

    def test_split_terms(self):
        self.assertEqual(split_terms("A; B;\nC"), ["A", "B", "C"])

    def test_filename_safe(self):
        self.assertEqual(safe_filename("NDA / Mutual?", "pdf", "2026-10-03"), "LegalEase_NDA_Mutual_2026-10-03.pdf")

    def test_html_preview_escapes(self):
        out = format_html_preview("# T\n<script>alert(1)</script>")
        self.assertNotIn("<script>", out)
        self.assertIn("&lt;script&gt;", out)


class ModelTests(unittest.TestCase):
    def test_empty_fields_rejected(self):
        from pydantic import ValidationError
        from backend.models import DocumentRequest

        for field in ("document_type", "parties", "terms", "effective_date"):
            data = dict(document_type="NDA", parties="A", terms="B", effective_date="1 Jan 2026")
            data[field] = "   "
            with self.assertRaises(ValidationError):
                DocumentRequest(**data)


class ExportTests(unittest.TestCase):
    def test_txt(self):
        data = generate_txt(sample_doc("NDA"))
        text = data.decode("utf-8")
        self.assertIn("IN WITNESS WHEREOF", text)
        self.assertIn("Generated with LegalEase AI", text)

    def test_docx_all_types_with_table_footer_logo(self):
        from docx import Document

        for kind in DOC_TYPES:
            raw = format_docx(sample_doc(kind), kind, TERMS, LOGO)
            doc = Document(io.BytesIO(raw))
            self.assertEqual(len(doc.tables), 1)
            self.assertEqual(len(doc.tables[0].rows), 4)  # header + 3 terms
            self.assertIn("Generated with LegalEase AI", doc.sections[0].footer.paragraphs[0].text)
            self.assertEqual(len(doc.inline_shapes), 1)  # logo

    def test_docx_without_logo_or_terms(self):
        from docx import Document

        raw = format_docx(sample_doc("NDA"), "NDA", None, None)
        self.assertEqual(len(Document(io.BytesIO(raw)).tables), 0)
        raw = format_docx(sample_doc("NDA"), "NDA", TERMS, "/nonexistent/logo.png")
        self.assertTrue(raw.startswith(b"PK"))

    @unittest.skipUnless(importlib.util.find_spec("fpdf"), "fpdf2 not installed")
    def test_pdf_all_types_and_long(self):
        from document.pdf_generator import format_pdf

        for kind in DOC_TYPES:
            raw = format_pdf(sample_doc(kind), kind, TERMS, LOGO)
            self.assertTrue(raw.startswith(b"%PDF"))
        long_raw = format_pdf(sample_doc("Service Agreement", 120), "Service Agreement", TERMS, LOGO)
        import re
        self.assertGreater(len(re.findall(rb"/Type\s*/Page(?!s)", long_raw)), 3)
        self.assertTrue(format_pdf(sample_doc("NDA"), "NDA", None, None).startswith(b"%PDF"))


if __name__ == "__main__":
    unittest.main()
